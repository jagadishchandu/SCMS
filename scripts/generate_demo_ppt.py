"""Generate the SCMS demo deck (PowerPoint).

Run:
    python3 /app/scripts/generate_demo_ppt.py
Output:
    /app/SCMS_Demo_Deck.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ───────── Brand palette ─────────
PRIMARY = RGBColor(0x4E, 0x5B, 0x3A)   # sage
SAND    = RGBColor(0xE6, 0xE1, 0xD6)   # sand
INK     = RGBColor(0x2F, 0x3E, 0x2E)   # near-black
INK_SOFT= RGBColor(0x5C, 0x66, 0x56)
BG      = RGBColor(0xFA, 0xF8, 0xF2)   # cream
ACCENT  = RGBColor(0xC4, 0x55, 0x3A)   # terracotta
SUCCESS = RGBColor(0x4F, 0x7A, 0x4F)
DANGER  = RGBColor(0xB0, 0x40, 0x40)
WARNING = RGBColor(0xC7, 0x9A, 0x3A)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# ───────── Layout ─────────
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


def _blank_slide(bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg_rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg_rect.fill.solid(); bg_rect.fill.fore_color.rgb = bg
    bg_rect.line.fill.background()
    return s


def _tb(slide, x, y, w, h, text, *, size=18, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font_name="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if not p.runs:
        r = p.add_run()
    else:
        r = p.runs[0]
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font_name
    return tb


def _bullets(slide, x, y, w, h, items, *, size=14, color=INK,
              indent_color=PRIMARY, font_name="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.05); tf.margin_bottom = Inches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        # Bullet glyph
        r1 = p.add_run(); r1.text = "▎ "
        r1.font.size = Pt(size); r1.font.color.rgb = indent_color; r1.font.bold = True; r1.font.name = font_name
        r2 = p.add_run(); r2.text = item
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = font_name
        p.space_after = Pt(6)
    return tb


def _accent_bar(slide, x, y, h=Inches(0.08), w=Inches(0.7), color=ACCENT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def _card(slide, x, y, w, h, fill=WHITE, line=SAND):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.04
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line
    c.line.width = Pt(0.75)
    c.shadow.inherit = False
    return c


def _section_header(slide, eyebrow, title, *, accent=ACCENT):
    _accent_bar(slide, Inches(0.6), Inches(0.55), color=accent)
    _tb(slide, Inches(0.6), Inches(0.7), Inches(12), Inches(0.4),
        eyebrow.upper(), size=11, bold=True, color=INK_SOFT)
    _tb(slide, Inches(0.6), Inches(1.05), Inches(12), Inches(1.0),
        title, size=34, bold=True, color=INK)


def _footer(slide, page_num, total, section):
    _tb(slide, Inches(0.6), Inches(7.05), Inches(6), Inches(0.3),
        f"SCMS — School Child Monitoring System  ·  {section}",
        size=9, color=INK_SOFT)
    _tb(slide, Inches(11.7), Inches(7.05), Inches(1.2), Inches(0.3),
        f"{page_num} / {total}", size=9, color=INK_SOFT, align=PP_ALIGN.RIGHT)


# ─────────────────────────────────────────────────────────────────────
# SLIDE BUILDERS
# ─────────────────────────────────────────────────────────────────────

def slide_title():
    s = _blank_slide(bg=PRIMARY)
    # Left band accent
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), SLIDE_H)
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT; band.line.fill.background()
    # Logo block
    logo = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(0.9), Inches(0.9))
    logo.adjustments[0] = 0.15
    logo.fill.solid(); logo.fill.fore_color.rgb = SAND; logo.line.fill.background()
    _tb(s, Inches(1.0), Inches(1.0), Inches(0.9), Inches(0.9),
        "S", size=48, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _tb(s, Inches(2.1), Inches(1.05), Inches(6), Inches(0.45),
        "SCMS", size=22, bold=True, color=WHITE)
    _tb(s, Inches(2.1), Inches(1.45), Inches(8), Inches(0.45),
        "SCHOOL CHILD MONITORING SYSTEM", size=11, bold=True, color=SAND)

    _tb(s, Inches(1.0), Inches(2.8), Inches(11), Inches(1.4),
        "One platform.\nMany schools.\nOne unified parent experience.",
        size=44, bold=True, color=WHITE)
    _tb(s, Inches(1.0), Inches(5.05), Inches(11), Inches(0.5),
        "Multi-tenant SaaS for K-12 schools — students, attendance, marks, fees, "
        "HR & payroll, WhatsApp/Telegram alerts in 8 languages.",
        size=15, color=SAND)

    _tb(s, Inches(1.0), Inches(6.4), Inches(8), Inches(0.4),
        "Product demo  ·  Version 1.18.1  ·  February 2026",
        size=11, bold=True, color=SAND)
    return s


def slide_tagline_strip():
    s = _blank_slide(bg=BG)
    _section_header(s, "Why SCMS", "Built for the realities of running a school.")
    cards = [
        ("Strict RBAC", "6 roles, every endpoint scoped. No teacher sees another teacher's class. No parent sees another child."),
        ("FREE first", "Meta WhatsApp Cloud API (1000 free/month) + Telegram Bot (truly free, unlimited) before falling back to paid Twilio."),
        ("8 Indian languages", "Templates baked into the backend. Per-parent override beats the school default."),
        ("Per-tenant DBs", "Each school's operational data physically lives in `scms_db_school_<code>` — clean billing, clean backups, clean GDPR boundary."),
        ("Production hardened", "283 backend pytests, audit log on every important state change, JWT brute-force throttle, idempotent seeds."),
        ("Docker one-liner", "`docker compose up --build` brings up Mongo + FastAPI + React with a seeded demo school."),
    ]
    col_w = Inches(4.0); row_h = Inches(1.55)
    cols = 3
    for i, (h, body) in enumerate(cards):
        col = i % cols; row = i // cols
        x = Inches(0.5) + col * (col_w + Inches(0.1))
        y = Inches(2.0) + row * (row_h + Inches(0.15))
        _card(s, x, y, col_w, row_h)
        _accent_bar(s, x + Inches(0.25), y + Inches(0.2))
        _tb(s, x + Inches(0.25), y + Inches(0.35), col_w - Inches(0.4), Inches(0.4),
            h, size=14, bold=True, color=INK)
        _tb(s, x + Inches(0.25), y + Inches(0.75), col_w - Inches(0.4), Inches(0.8),
            body, size=10, color=INK_SOFT)
    return s


def slide_roles():
    s = _blank_slide(bg=BG)
    _section_header(s, "Roles & Access", "Six personas. Zero overlap.")
    roles = [
        ("Platform Super Admin", PRIMARY, [
            "Onboards new schools",
            "Approves subscription requests",
            "Compliance + audit",
        ]),
        ("School Super Admin", ACCENT, [
            "Full school control",
            "HR, payroll, finance",
            "Subscription + branding",
        ]),
        ("Teacher / Class Teacher", SUCCESS, [
            "Attendance + marks for own classes",
            "Reimbursement requests",
            "Cannot see other classes",
        ]),
        ("Accounts", WARNING, [
            "Fees, payments, expenses",
            "View payroll PDF",
            "No academic data",
        ]),
        ("Parent", INK_SOFT, [
            "Only own child's record",
            "WhatsApp / Telegram digests",
            "Pay fees, download receipts",
        ]),
        ("Student (planned)", DANGER, [
            "Analytics on own record",
            "Backlog item — coming soon",
        ]),
    ]
    col_w = Inches(4.0); row_h = Inches(2.05)
    cols = 3
    for i, (h, col_c, lines) in enumerate(roles):
        col = i % cols; row = i // cols
        x = Inches(0.5) + col * (col_w + Inches(0.1))
        y = Inches(2.0) + row * (row_h + Inches(0.18))
        _card(s, x, y, col_w, row_h)
        # color bar on left of card
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), row_h)
        bar.fill.solid(); bar.fill.fore_color.rgb = col_c; bar.line.fill.background()
        _tb(s, x + Inches(0.35), y + Inches(0.2), col_w - Inches(0.5), Inches(0.4),
            h, size=14, bold=True, color=col_c)
        _bullets(s, x + Inches(0.25), y + Inches(0.7), col_w - Inches(0.4), row_h - Inches(0.8),
                 lines, size=11, color=INK_SOFT, indent_color=col_c)
    return s


def slide_architecture():
    s = _blank_slide(bg=BG)
    _section_header(s, "Architecture", "React + FastAPI + Per-tenant MongoDB.")

    boxes = [
        ("Frontend",  Inches(0.6),  Inches(2.0),  Inches(3.6), Inches(2.5),
         ["React 18 + Tailwind", "Phosphor icons", "Recharts", "Axios (httpOnly cookie)", "Role-scoped routing"]),
        ("Backend",   Inches(4.85), Inches(2.0),  Inches(3.6), Inches(2.5),
         ["FastAPI + Motor", "JWT (bcrypt + brute-force)", "ReportLab PDFs", "APScheduler crons", "29+ route modules"]),
        ("Data layer", Inches(9.1), Inches(2.0),  Inches(3.6), Inches(2.5),
         ["Master DB: scms_db", "Per-school DB: scms_db_school_<code>", "31 tenant collections",
          "Audit dual-write", "Idempotent migrations"]),
    ]
    for label, x, y, w, h, items in boxes:
        _card(s, x, y, w, h)
        _accent_bar(s, x + Inches(0.25), y + Inches(0.2))
        _tb(s, x + Inches(0.25), y + Inches(0.35), w - Inches(0.4), Inches(0.4),
            label, size=15, bold=True, color=INK)
        _bullets(s, x + Inches(0.2), y + Inches(0.85), w - Inches(0.3), h - Inches(1.0),
                 items, size=11, color=INK_SOFT)
    # bottom strip
    _card(s, Inches(0.6), Inches(4.85), Inches(12.1), Inches(1.5))
    _accent_bar(s, Inches(0.85), Inches(5.05))
    _tb(s, Inches(0.85), Inches(5.2), Inches(11.5), Inches(0.4),
        "External services (opt-in, all free-first)", size=14, bold=True, color=INK)
    _tb(s, Inches(0.85), Inches(5.6), Inches(11.5), Inches(0.7),
        "Meta WhatsApp Cloud · Twilio fallback · Telegram Bot · SMTP (Gmail/SES/SendGrid/Mailgun) · "
        "Azure Blob (planned) · Stripe/Razorpay (planned)",
        size=11, color=INK_SOFT)
    return s


def feature_slide(eyebrow, title, value_line, bullets, screen_hint, accent=ACCENT):
    s = _blank_slide(bg=BG)
    _section_header(s, eyebrow, title, accent=accent)
    # Value strap
    strap = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(2.05), Inches(12.1), Inches(0.65))
    strap.adjustments[0] = 0.2
    strap.fill.solid(); strap.fill.fore_color.rgb = SAND; strap.line.fill.background()
    _tb(s, Inches(0.85), Inches(2.05), Inches(11.6), Inches(0.65),
        value_line, size=14, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    # Bullets (left)
    _tb(s, Inches(0.6), Inches(3.0), Inches(7.5), Inches(0.4),
        "What it does", size=13, bold=True, color=accent)
    _bullets(s, Inches(0.6), Inches(3.4), Inches(7.5), Inches(3.5),
             bullets, size=12, color=INK, indent_color=accent)
    # Right-side card describing where to demo it
    _card(s, Inches(8.4), Inches(3.0), Inches(4.3), Inches(3.9))
    _accent_bar(s, Inches(8.65), Inches(3.2), color=accent)
    _tb(s, Inches(8.65), Inches(3.35), Inches(3.9), Inches(0.4),
        "Demo this in", size=13, bold=True, color=INK)
    _bullets(s, Inches(8.55), Inches(3.85), Inches(4.0), Inches(2.9),
             screen_hint, size=11, color=INK_SOFT, indent_color=accent)
    return s


# ─── Per-feature content ────────────────────────────────────────────

FEATURES = [
    # (eyebrow, title, strap, bullets, demo-locations, accent)
    ("Onboarding", "Single Super Admin · Tenant per school",
     "One Platform Super Admin (env-driven). One School Super Admin per school. Zero ambiguity.",
     [
         "Platform Super Admin is provisioned from ADMIN_EMAIL / ADMIN_PASSWORD on first boot.",
         "Creating a school auto-creates a dedicated MongoDB DB `scms_db_school_<code>` with 31 collections.",
         "Creating a school auto-creates exactly ONE school admin (cannot be created elsewhere).",
         "Single-admin policy is enforced at startup: extra `platform_admin` rows are auto-deactivated.",
     ],
     [
         "Login as Platform Super Admin → /platform → 'Add school'",
         "Fill school name, code, admin name/email/password",
         "Switch to the new school admin to verify auto-login works",
     ], PRIMARY),

    ("Authentication", "JWT cookies · Brute-force throttle · Wrong-role guard",
     "Strict role-validation on login: even Platform Super Admin must pick the right role in the dropdown.",
     [
         "httpOnly JWT cookie + refresh token; bcrypt password hashes.",
         "5 failed attempts in 15 min → 15-min lockout per (ip, email).",
         "Picking the wrong role at login returns an error AND clears the cookie (no back-door via direct URL).",
         "Generic error messages — no email enumeration.",
     ],
     [
         "Login as Teacher with role=Parent → blocked toast",
         "Then login with role=Teacher → /teacher dashboard",
         "5 wrong passwords → 429 Too Many Attempts",
     ], PRIMARY),

    ("Password Reset", "Telegram OTP — no email required",
     "Survives lost passwords without email. 6-digit CSPRNG OTP, bcrypt-hashed, 10-min TTL.",
     [
         "User clicks 'Forgot password?' on /login.",
         "Backend generates 6-digit OTP, stores ONLY its bcrypt hash, sends plaintext to user's linked Telegram chat.",
         "Verify with email + code + new password. 5 wrong attempts → reset locked.",
         "30-sec cooldown between requests. Generic response — no email enumeration.",
     ],
     [
         "/login → Forgot password?",
         "Enter user email → OTP arrives on Telegram",
         "Enter OTP + new password → redirect to /login",
     ], PRIMARY),

    ("School Admin", "Students · Teachers · Classes · Attendance",
     "Everything a principal needs in one console — strictly scoped to their school.",
     [
         "Students: name, photo, parent contact, class, section, UID — manual or Excel bulk upload.",
         "Teachers: subject mapping required; Class Teacher promotion is one click.",
         "Classes & Subjects: school admin can delete; teachers can ADD but not remove.",
         "Attendance: traditional grid OR teacher mobile swipe-card flow.",
     ],
     [
         "/school → 'View roster' drawer (live tenant counts)",
         "/school/students → add a student",
         "/school/classes → add a section as teacher (no delete buttons visible)",
     ], ACCENT),

    ("Exams & Marks", "Per-subject marks · Publish triggers parent alerts",
     "Teachers enter marks per exam; publishing fires a notification per parent in their preferred language.",
     [
         "Exams have date, max marks, type (unit / mid / final).",
         "Marks entry locked to the teacher's assigned subject for that class.",
         "Publish → WhatsApp + Email + Telegram (configured channels) per parent.",
         "Audit log records who published which exam when.",
     ],
     [
         "/school/exams → create an exam",
         "/teacher/marks → enter marks",
         "Click Publish → check /school/notifications log",
     ], ACCENT),

    ("Attendance", "Mobile swipe cards · Auto WhatsApp on absent",
     "3 large status buttons per student (Present · Absent · Not in class). Auto parent alert on absent.",
     [
         "Phase 1: pick class/section/date. Phase 2: card-by-card with progress bar. Phase 3: review & submit.",
         "Status options: present · absent · sick · not_in_class.",
         "When status=absent, a notification is queued for the parent (WhatsApp + Telegram).",
         "Sick-leave entry auto-marks attendance=sick for weekdays in the span (no double entry).",
     ],
     [
         "/teacher/attendance-cards → walk through 3 phases on a phone",
         "Mark one student absent → /parent/calendar shows it within seconds",
     ], ACCENT),

    ("Fees", "Per-class fee structures · PDF receipts",
     "Accounts staff record payments; PDF receipts download from parent portal & analytics page.",
     [
         "Fee structures per class (tuition, transport, books, etc.).",
         "Payments tracked with mode (cash/UPI/card/transfer), reference, partial/paid/pending status.",
         "Each payment generates a ReportLab PDF receipt at /api/receipts/payment/{id}.pdf.",
         "Per-student academic+fees report PDF for parent download.",
     ],
     [
         "/accounts → record a payment for STU-00001",
         "Parent portal → Download receipt",
         "Student Analytics → Download student report",
     ], ACCENT),

    ("Sick Leaves", "Multi-day · Medical cert upload · Recurring detection",
     "Recurring-illness widget flags students with ≥3 sick leaves in 90 days.",
     [
         "School admin / teacher (own class) / parent (own child) can record.",
         "Cert upload: PDF/image up to 5 MB (data URL in tenant DB).",
         "Auto-attendance: weekdays in the span are marked `sick` (no double entry).",
         "Recurring view groups students by leave count and total days off.",
     ],
     [
         "/school/sick-leaves → record 5-day flu for STU-00001",
         "Upload certificate → status badge updates",
         "Recurring section shows the top offenders",
     ], ACCENT),

    ("Analytics", "Risk flags · Donut + line + bar charts",
     "Predictive scoring surfaces 'students at risk' on the school admin dashboard.",
     [
         "Rolling 30-day attendance + marks drop + absolute floor penalties.",
         "Top 5 at-risk students card on /school dashboard with reasons (badges) + click-through.",
         "Student analytics page: attendance donut, marks trend line, fees stacked bar, exam-wise bars.",
         "Parent portal mirrors the same charts for their child.",
     ],
     [
         "/school dashboard → 'Students at risk' card",
         "Click into a student → full chart page",
         "/parent → same charts for own child only",
     ], ACCENT),

    ("Calendar", "Color-coded holidays + exams · 60-day upcoming list",
     "Monthly calendar UI for school admin, teachers, and parents — all roles share the same color coding.",
     [
         "Backend `/api/calendar/events` returns holidays + exams in one call.",
         "Prev / Next / Today nav. Click a date to see all events.",
         "Upcoming-60-days list under the grid for at-a-glance planning.",
     ],
     [
         "/school/calendar → flip months",
         "/teacher/calendar → same view scoped to their classes",
         "/parent/calendar → same for parent's school",
     ], ACCENT),

    ("WhatsApp / Telegram / Email", "Free-first multi-channel notifications",
     "Meta Cloud API (1000 free/mo) → Twilio fallback → SMTP email → Telegram bot. All multilingual.",
     [
         "8 Indian languages baked in (en, hi, ta, te, kn, mr, bn, gu). Parent override beats school default.",
         "Templates: attendance_absent · marks_published · payment_confirmed · monthly_digest · subscription_renewal.",
         "Webhook handlers ingest Meta delivery statuses + Telegram /start linking.",
         "When no provider is configured, notifications stay `queued` with a clear `delivery_note`.",
     ],
     [
         "/school/settings → 'Send test WhatsApp / Telegram / Email'",
         "/school/notifications → log with provider + status per row",
         "Switch language in /school/settings → next alert is in that language",
     ], SUCCESS),

    ("Subscription Billing", "₹30/child/month · ₹300/child/year",
     "Per-child pricing. School admin submits payment, Platform admin confirms.",
     [
         "Live quote computed from tenant DB's active student count.",
         "Monthly = ₹30/child/month · Yearly = ₹300/child/year (≈16% discount).",
         "Workflow: school admin submits payment → status `pending_confirmation` → platform admin Confirm → `subscription` becomes active.",
         "Confirming extends validity from `valid_until` if currently active (so a renewal mid-cycle isn't punished).",
     ],
     [
         "/school/subscription → submit a monthly payment",
         "/platform/subscriptions → Confirm",
         "Check /school/subscription history table",
     ], WARNING),

    ("Subscription Expiry Guard", "Auto-block writes after grace · Renewal nudge",
     "402 + structured body when expired. Site-wide banner. 7-day Telegram/WhatsApp nudge.",
     [
         "Middleware blocks POST/PATCH/PUT/DELETE for schools where `valid_until + grace_days < now`.",
         "/api/auth/* and /api/subscriptions/* always stay open so users can pay.",
         "SubscriptionBanner: amber when ≤7 days, red in grace, red+blocked when fully expired.",
         "Cron job at 08:00 IST queues a renewal nudge for any school inside the 7-day window (24h cooldown).",
     ],
     [
         "platform/subscriptions → set valid_until to T-1 day for a school",
         "Login as that school's admin → red banner + 402 on any write",
         "/api/subscriptions/nudge-renewals → manual trigger",
     ], WARNING),

    ("v1.18 · Approval Workflow", "Lockdown until Platform Super Admin approves",
     "New schools must request approval with a children-count. Until approved, the entire school is HTTP 423 locked except Dashboard + Subscription.",
     [
         "school.subscription_approval.status ∈ {not_started · pending · approved · rejected}.",
         "Sidebar shows ONLY Dashboard + Subscription for non-approved schools.",
         "Allowed endpoints during lockdown: /auth/*, /subscriptions/*, /analytics/school, /people, /db-status.",
         "Legacy schools were auto-migrated to `approved` on startup so nothing breaks.",
     ],
     [
         "Platform admin creates a brand new school (no auto_approve)",
         "School admin logs in → red banner + locked sidebar",
         "Submits requested_children_count = 200 → status → pending",
         "Platform admin /platform/subscriptions → Approve → school unlocked",
     ], DANGER),

    ("v1.18.1 · Editable Count + Compliance",
     "School admin can edit at any time · Platform admin sees a separate read-only Compliance view",
     "Post-approval, the school admin can update their requested count without re-approval — surfaces in the Compliance dashboard.",
     [
         "POST /api/subscriptions/approval-request is status-aware: approved → just updates count (status preserved); else → resubmits as pending.",
         "/platform/compliance: live actual student count vs requested/approved with `over_capacity` flag and variance.",
         "5 filter pills + 5 KPI tiles. STRICTLY view-only — no Approve/Reject buttons here.",
         "Used for billing reconciliation and capacity planning.",
     ],
     [
         "Login as principal@sunrise.edu → /school/subscription → Edit count",
         "Change 156 → 700, save → status still approved",
         "Switch to platform admin → /platform/compliance → see new requested value",
         "Filter `over_capacity` to see only schools above their cap",
     ], DANGER),

    ("v1.17 · HR — Non-Teaching Staff", "Cleaners · Drivers · Security · Cooks · Coaches",
     "Track every paid person on campus — not just teachers — so payroll is one source of truth.",
     [
         "10 role types: cleaner, bus_driver, security, cook, sports_coach, nurse, librarian, lab_assistant, office_assistant, other.",
         "Each row stores monthly_salary, phone, email, joining_date, notes.",
         "Soft delete (is_active=false) — keeps history for payroll audits.",
         "Combined payroll summary: teachers + accounts + non-teaching = monthly + yearly outflow.",
     ],
     [
         "/school/staff → 3 KPI tiles (active staff / monthly / yearly)",
         "Add a sports_coach with ₹25,000/mo",
         "Edit, then soft-delete → no longer in active roster",
     ], SUCCESS),

    ("v1.17 · Finance — Expenses",
     "Direct expenses · Teacher reimbursement requests",
     "School admin records direct expenses (auto-approved); teachers submit out-of-pocket requests for approval.",
     [
         "Categories: salary · supplies · maintenance · utilities · travel · events · transport · food · stationery · books · other.",
         "Teacher submission → status=pending → admin/accounts decides (approve / reject / mark_paid).",
         "Flag `reimburse_with_salary=true` so it shows up in the next monthly payroll PDF.",
         "Receipt upload (data URL) attached per expense.",
     ],
     [
         "/teacher/expenses → submit ₹1,250 'travel' with `reimburse_with_salary`",
         "/school/expenses → approve it",
         "/school/staff → Monthly Payroll PDF includes the reimbursement",
     ], SUCCESS),

    ("v1.18 · Monthly Payroll PDF",
     "One-click ReportLab PDF — teachers + accounts + staff + reimbursements",
     "Hand a single document to your accountant. Base salary + approved reimbursements + per-section subtotals + grand total.",
     [
         "GET /api/staff/payroll/monthly-pdf?month=YYYY-MM (defaults to current month).",
         "Sections: Teaching staff · Accounts staff · Non-teaching support staff.",
         "Each row: # · Name · Role/Contact · Base Salary · Reimbursements · Net Payable.",
         "Subtotals per section + branded Grand Total band at the bottom.",
         "RBAC: school_admin + accounts only (teachers 403). Invalid month format → 400.",
     ],
     [
         "/school/staff → 'Monthly Payroll PDF' button",
         "Enter 2026-02 → PDF downloads",
         "Open PDF → verify subtotals + grand total + reimbursable expenses line",
     ], SUCCESS),

    ("Bulk Excel Uploads", "Students · Attendance · Holidays",
     "School admin + teachers (own classes) can upload .xlsx — wrong-section rows are reported, not blocking.",
     [
         "Template download for each entity type — same shape as the API.",
         "Server-side validation: missing fields per row → goes into `errors[]`, valid rows still imported.",
         "Teachers uploading students get a per-row error for any class they don't teach.",
         "Audit log records the upload + the count of imported/failed rows.",
     ],
     [
         "/school/uploads → download Students template",
         "Upload a file with 5 valid + 2 invalid rows",
         "Result: 5 imported, 2 in errors[] with reason",
     ], ACCENT),

    ("Audit Log", "Filter + Search + CSV Export",
     "Every important state change has a row — filter by role/action/resource/date, download UTF-8 CSV with BOM.",
     [
         "Captures user create/update/delete, student photo upload, exam publish, payment confirm, approval decisions, payroll PDF generation, etc.",
         "School admin sees only their school's events; platform admin sees everything.",
         "Filters: actor_role · action · resource · school_id · date_from/to · free-text · limit.",
         "CSV export keeps filters and opens cleanly in Excel (UTF-8 BOM).",
     ],
     [
         "/school/audit → filter `action=publish_exam, last 30 days`",
         "Click 'Download CSV' → Excel opens",
         "/platform/audit → same UI with cross-school visibility",
     ], INK_SOFT),

    ("Mobile Responsive", "Sidebar collapses · Parent bottom-nav",
     "Built for phones first. Teacher swipe-card attendance shines on a 5-inch screen.",
     [
         "Sidebar collapses to a slide-in drawer below 1024px.",
         "Parent Portal has a dedicated bottom navigation bar.",
         "All forms + tables use responsive 1/2/3/4-column grids.",
         "Tested at 390×844 (iPhone 14) — drawer opens/closes correctly.",
     ],
     [
         "Resize browser to 400×900",
         "Hamburger menu → drawer slides in",
         "/teacher/attendance-cards → big touch buttons",
     ], ACCENT),

    ("Branding & Dark Mode", "Per-school logo / color / tagline · localStorage dark toggle",
     "Each school's parent portal feels like THEIR school, not 'SCMS'.",
     [
         "/school/settings: upload logo (data URL), set primary color hex, set tagline.",
         "Applied across school admin, teacher, accounts, and parent pages.",
         "Dark mode toggle in the top bar; persisted to localStorage.",
         "Future: per-school custom domain via reverse-proxy rule.",
     ],
     [
         "/school/settings → upload a logo + set primary color",
         "Login as parent → see new logo + color",
         "Top bar → toggle dark mode",
     ], PRIMARY),
]


def slide_test_credentials():
    s = _blank_slide(bg=BG)
    _section_header(s, "Hands-on demo", "Default credentials (seeded automatically)", accent=PRIMARY)
    headers = ["Role", "Email", "Password"]
    rows = [
        ("Platform Super Admin", "jagadishchandu789@gmail.com", "Manasa@123"),
        ("School Super Admin (DEMO)", "principal@sunrise.edu", "School@123"),
        ("Teacher (Class Teacher)", "priya@sunrise.edu", "Teacher@123"),
        ("Teacher", "rohit@sunrise.edu", "Teacher@123"),
        ("Accounts", "accounts@sunrise.edu", "Accounts@123"),
        ("Parent of STU-00001", "parent.stu-00001@sunrise.edu", "Parent@123"),
    ]
    # Table
    top = Inches(2.1); left = Inches(0.6); col_w = [Inches(3.6), Inches(5.4), Inches(3.1)]
    # Header
    hx = left
    for i, h in enumerate(headers):
        cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, hx, top, col_w[i], Inches(0.5))
        cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY; cell.line.fill.background()
        _tb(s, hx + Inches(0.15), top, col_w[i] - Inches(0.2), Inches(0.5),
            h.upper(), size=11, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        hx += col_w[i]
    # Body
    ry = top + Inches(0.5)
    for r_idx, row in enumerate(rows):
        rx = left
        bg = WHITE if r_idx % 2 == 0 else SAND
        for i, val in enumerate(row):
            cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, ry, col_w[i], Inches(0.5))
            cell.fill.solid(); cell.fill.fore_color.rgb = bg; cell.line.color.rgb = SAND
            cell.line.width = Pt(0.5)
            _tb(s, rx + Inches(0.15), ry, col_w[i] - Inches(0.2), Inches(0.5),
                val, size=12, color=INK, anchor=MSO_ANCHOR.MIDDLE,
                font_name=("Consolas" if i > 0 else "Calibri"))
            rx += col_w[i]
        ry += Inches(0.5)

    # Footer hint
    _tb(s, Inches(0.6), Inches(6.0), Inches(12), Inches(0.5),
        "Tip: the login form requires explicit role selection from the dropdown — wrong-role attempts are rejected even with valid creds.",
        size=11, color=INK_SOFT)
    _tb(s, Inches(0.6), Inches(6.4), Inches(12), Inches(0.5),
        "Forgot password? Click 'Forgot password?' on the login page → 6-digit OTP delivered to the user's Telegram chat.",
        size=11, color=INK_SOFT)
    return s


def slide_test_status():
    s = _blank_slide(bg=BG)
    _section_header(s, "Quality", "Tested. Hardened. Production-ready.", accent=SUCCESS)
    cards = [
        ("283", "Backend pytest cases",
         "Auth · RBAC · CRUD · analytics · subscriptions · audit · PDF receipts · Twilio scaffold · "
         "WhatsApp / Telegram / SMTP · Multilingual · Sick leaves · Approval workflow · Payroll PDF · Compliance."),
        ("31",  "Tenant collections per school",
         "Idempotent schema initialiser. `/api/schools/{id}/db-status` verifies missing collections in real time."),
        ("8",   "Indian languages",
         "5 templates × 8 languages = 40 baked-in entries. Parent override beats school default."),
        ("423", "HTTP code on approval lockdown",
         "Structured body: `{code: approval_pending|approval_required, renew_url}`. Sidebar dynamically hides nav items."),
    ]
    col_w = Inches(6.05); row_h = Inches(2.0)
    for i, (big, title, body) in enumerate(cards):
        col = i % 2; row = i // 2
        x = Inches(0.5) + col * (col_w + Inches(0.15))
        y = Inches(2.1) + row * (row_h + Inches(0.2))
        _card(s, x, y, col_w, row_h)
        _accent_bar(s, x + Inches(0.3), y + Inches(0.3), color=SUCCESS)
        _tb(s, x + Inches(0.3), y + Inches(0.45), Inches(2.2), Inches(0.9),
            big, size=46, bold=True, color=SUCCESS)
        _tb(s, x + Inches(2.5), y + Inches(0.45), col_w - Inches(2.7), Inches(0.4),
            title, size=14, bold=True, color=INK)
        _tb(s, x + Inches(2.5), y + Inches(0.85), col_w - Inches(2.7), row_h - Inches(0.95),
            body, size=11, color=INK_SOFT)
    return s


def slide_roadmap():
    s = _blank_slide(bg=BG)
    _section_header(s, "What's next", "Roadmap (post v1.18.1)", accent=ACCENT)
    cols = [
        ("P1 — next up", PRIMARY, [
            "Stripe / Razorpay automated payments",
            "Alert when school crosses `over_capacity`",
            "'Send test email' on Settings",
        ]),
        ("P2 — backlog", ACCENT, [
            "Azure Blob for uploads (photos / certs / PDFs)",
            "Recovery nudges post-expiry (day 1 / 3 / 7)",
            "Auto-deduct reimbursements from next payroll",
            "Per-school WhatsApp template editor",
        ]),
        ("P3 — future", INK_SOFT, [
            "React Native mobile shell",
            "Public school directory + admissions inquiry",
            "Voice-to-text attendance",
            "Student login & analytics portal",
        ]),
    ]
    col_w = Inches(4.0); top = Inches(2.1)
    for i, (title, color, items) in enumerate(cols):
        x = Inches(0.5) + i * (col_w + Inches(0.1))
        _card(s, x, top, col_w, Inches(4.5))
        _accent_bar(s, x + Inches(0.3), top + Inches(0.25), color=color)
        _tb(s, x + Inches(0.3), top + Inches(0.4), col_w - Inches(0.4), Inches(0.5),
            title, size=15, bold=True, color=color)
        _bullets(s, x + Inches(0.25), top + Inches(0.95), col_w - Inches(0.4), Inches(3.4),
                 items, size=12, color=INK, indent_color=color)
    # Footer line
    _tb(s, Inches(0.6), Inches(6.8), Inches(12), Inches(0.4),
        "Recently shipped: v1.18.1 (editable count + compliance) · v1.18 (approval workflow + payroll PDF) · v1.17 (HR/Finance) · v1.16 (platform drawer) · v1.15 (per-tenant DB authority)",
        size=10, color=INK_SOFT)
    return s


def slide_closing():
    s = _blank_slide(bg=PRIMARY)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), SLIDE_H)
    band.fill.solid(); band.fill.fore_color.rgb = ACCENT; band.line.fill.background()
    _tb(s, Inches(1.0), Inches(2.4), Inches(11), Inches(1.5),
        "Run a real school.\nNot a spreadsheet farm.",
        size=48, bold=True, color=WHITE)
    _tb(s, Inches(1.0), Inches(4.6), Inches(11), Inches(0.6),
        "Set up in 1 docker command. Onboard a school in 90 seconds.\nTwo gates, FREE WhatsApp, zero lock-in.",
        size=16, color=SAND)
    _tb(s, Inches(1.0), Inches(6.4), Inches(11), Inches(0.4),
        "github.com/your-org/scms  ·  Built with FastAPI + React + MongoDB",
        size=11, color=SAND)
    return s


# ─────────────────────────────────────────────────────────────────────
# Build the deck
# ─────────────────────────────────────────────────────────────────────

def build():
    # Title + intro
    slide_title()
    slide_tagline_strip()
    slide_roles()
    slide_architecture()

    # Per-feature deep-dives
    for f in FEATURES:
        feature_slide(*f)

    # Wrap up
    slide_test_credentials()
    slide_test_status()
    slide_roadmap()
    slide_closing()

    # Apply page-of-pages footer after everything is built
    total = len(prs.slides)
    sections = [
        "Overview", "Overview", "Overview", "Architecture",
    ] + [f"Feature {i+1}" for i in range(len(FEATURES))] + [
        "Demo creds", "Quality", "Roadmap", "Wrap"
    ]
    for idx, slide in enumerate(prs.slides):
        if idx == 0 or idx == total - 1:
            continue  # skip title / closing splash
        section = sections[idx] if idx < len(sections) else ""
        _footer(slide, idx + 1, total, section)

    out_path = Path("/app/SCMS_Demo_Deck.pptx")
    prs.save(out_path)
    print(f"Wrote {out_path}  ({total} slides)")


if __name__ == "__main__":
    build()
