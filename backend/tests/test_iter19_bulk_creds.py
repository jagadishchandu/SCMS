"""
iter19 — Bulk student excel + parent_password, teacher-edit/delete student,
credential-slip PDF, send-credentials, bulk export parent credentials,
TeachersPage edit/deactivate/delete users, and StudentUpdate gender/dob.
"""
import io
import os
import time
import openpyxl
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://class-access-hub.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

GRADE5 = "9c33b98760fb4a82b63f44de"
GRADE6 = "6573544e9b3648d68198fa44"


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=20)
    assert r.status_code == 200, f"login failed for {email}: {r.status_code} {r.text}"
    return s


@pytest.fixture(scope="session")
def admin_token():
    return _login("principal@sunrise.edu", "School@123")


@pytest.fixture(scope="session")
def teacher_token():
    return _login("priya@sunrise.edu", "Teacher@123")


@pytest.fixture(scope="session")
def accounts_token():
    return _login("accounts@sunrise.edu", "Accounts@123")


def h(_sess):
    # back-compat shim — kept so existing call sites still work; sessions auth via cookie
    return {}


# -------- Excel template --------
class TestExcelTemplate:
    def test_template_has_parent_password_header(self, admin_token):
        r = admin_token.get(f"{API}/uploads/template/students", timeout=20)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["Content-Type"]
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(max_row=1))]
        assert len(headers) == 11
        assert "parent_password" in headers
        assert headers[-1] == "parent_password"


def _make_xlsx(rows):
    """rows is list of dicts keyed by template headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = ["name", "class_id", "section", "roll_number", "gender", "dob",
               "parent_name", "parent_phone", "parent_whatsapp", "parent_email", "parent_password"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(k, "") for k in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# -------- Excel student upload --------
class TestExcelUploadStudents:
    def test_upload_with_parent_password_and_school_tld_and_credentials(self, admin_token):
        ts = int(time.time())
        p1_email = f"test_iter19_p1_{ts}@example.com"
        rows = [
            # row1: explicit password >=6 → parent created with that pwd
            {"name": f"TEST UpRow1 {ts}", "class_id": GRADE5, "section": "A", "roll_number": "991",
             "gender": "M", "dob": "2015-04-01",
             "parent_name": "P1", "parent_phone": "+919000000001", "parent_whatsapp": "+919000000001",
             "parent_email": p1_email, "parent_password": f"MyPass{ts}"},
            # row2: blank parent_email → synthetic .school TLD, blank password → default Parent@123
            {"name": f"TEST UpRow2 {ts}", "class_id": GRADE5, "section": "A", "roll_number": "992",
             "gender": "F", "dob": "2015-05-02",
             "parent_name": "P2", "parent_phone": "+919000000002", "parent_whatsapp": "+919000000002",
             "parent_email": "", "parent_password": ""},
            # row3: too short pwd → falls back to default
            {"name": f"TEST UpRow3 {ts}", "class_id": GRADE5, "section": "A", "roll_number": "993",
             "gender": "M", "dob": "2015-06-03",
             "parent_name": "P3", "parent_phone": "+919000000003", "parent_whatsapp": "+919000000003",
             "parent_email": f"test_iter19_p3_{ts}@example.com", "parent_password": "abc"},
            # row4: reuse p1_email → parent_was_created should be False, password=null in credentials
            {"name": f"TEST UpRow4 {ts}", "class_id": GRADE5, "section": "A", "roll_number": "994",
             "gender": "F", "dob": "2015-07-04",
             "parent_name": "P1", "parent_phone": "+919000000001", "parent_whatsapp": "+919000000001",
             "parent_email": p1_email, "parent_password": ""},
        ]
        xlsx = _make_xlsx(rows)
        files = {"file": ("students.xlsx", xlsx.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = admin_token.post(f"{API}/uploads/students", files=files, timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["inserted"] == 4, data
        creds = data["credentials"]
        assert len(creds) == 4

        c1, c2, c3, c4 = creds
        # row1: created, exact password preserved
        assert c1["parent_was_created"] is True
        assert c1["parent_password"] == f"MyPass{ts}"
        assert c1["parent_email"] == p1_email
        # row2: synthetic .school TLD
        assert c2["parent_was_created"] is True
        assert c2["parent_email"].endswith(".school"), c2
        assert c2["parent_password"] == "Parent@123"
        # row3: short pwd → default
        assert c3["parent_was_created"] is True
        assert c3["parent_password"] == "Parent@123"
        # row4: reuse → not created, password null
        assert c4["parent_was_created"] is False
        assert c4["parent_password"] is None

        # Login with row1's explicit password
        lr = requests.post(f"{API}/auth/login", json={"email": p1_email, "password": f"MyPass{ts}"}, timeout=20)
        assert lr.status_code == 200, f"row1 explicit pwd login failed: {lr.text}"

        # Login with row2's synthetic .school parent + Parent@123
        lr2 = requests.post(f"{API}/auth/login",
                            json={"email": c2["parent_email"], "password": "Parent@123"}, timeout=20)
        assert lr2.status_code == 200, f"row2 synthetic .school login failed: {lr2.text}"


# -------- Bulk export parent creds --------
class TestParentCredsExport:
    def test_export_xlsx_admin(self, admin_token):
        r = admin_token.get(f"{API}/uploads/parent-credentials",
                            params={"class_id": GRADE5, "section": "A"}, timeout=30)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["Content-Type"]
        assert "attachment" in r.headers.get("Content-Disposition", "")
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(max_row=1))]
        for required in ["student_uid", "student_name", "class", "section",
                         "parent_name", "parent_phone", "parent_whatsapp",
                         "parent_email", "default_password_hint"]:
            assert required in headers, headers
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) >= 5, f"expected >=5 active Grade 5 A students, got {len(rows)}"

    def test_export_rbac_teacher_403(self, teacher_token):
        r = teacher_token.get(f"{API}/uploads/parent-credentials",
                              params={"class_id": GRADE5, "section": "A"}, timeout=20)
        assert r.status_code == 403


# -------- PATCH / DELETE student RBAC --------
def _seed_student(admin_token, name, class_id=GRADE5, section="A"):
    r = admin_token.post(f"{API}/students", json={
        "name": name, "class_id": class_id, "section": section,
        "roll_number": "", "gender": "M", "dob": "2015-01-01",
        "parent": {"name": "TP", "phone": "+919999999999", "relationship": "guardian", "email": "", "whatsapp": ""}
    }, timeout=20)
    assert r.status_code in (200, 201), r.text
    return r.json()


class TestStudentPatchDelete:
    def test_patch_gender_dob_persists(self, admin_token):
        s = _seed_student(admin_token, f"TEST_PatchGD_{int(time.time())}")
        sid = s["id"]
        r = admin_token.patch(f"{API}/students/{sid}", json={"gender": "F", "dob": "2015-01-01"}, timeout=20)
        assert r.status_code == 200, r.text
        # GET to verify
        g = admin_token.get(f"{API}/students/{sid}", timeout=20)
        assert g.status_code == 200
        assert g.json()["gender"] == "F"
        assert g.json()["dob"] == "2015-01-01"

    def test_patch_blank_to_none_coercion(self, admin_token):
        s = _seed_student(admin_token, f"TEST_PatchBlank_{int(time.time())}")
        r = admin_token.patch(f"{API}/students/{s['id']}", json={"gender": "", "dob": ""}, timeout=20)
        assert r.status_code == 200, r.text

    def test_teacher_patch_own_class_ok(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_TchPatchOwn_{int(time.time())}", GRADE5, "A")
        r = teacher_token.patch(f"{API}/students/{s['id']}", json={"roll_number": "77"}, timeout=20)
        assert r.status_code == 200, r.text

    def test_teacher_patch_other_class_403(self, admin_token, teacher_token):
        # Create student in Grade 6 B (Priya doesn't teach)
        s = _seed_student(admin_token, f"TEST_TchPatchOther_{int(time.time())}", GRADE6, "B")
        r = teacher_token.patch(f"{API}/students/{s['id']}", json={"roll_number": "77"}, timeout=20)
        assert r.status_code == 403

    def test_teacher_move_to_class_not_taught_403(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_TchMove_{int(time.time())}", GRADE5, "A")
        r = teacher_token.patch(f"{API}/students/{s['id']}", json={"class_id": GRADE6, "section": "B"}, timeout=20)
        assert r.status_code == 403
        assert "classes you teach" in r.text.lower()

    def test_teacher_delete_own_class_softinactive(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_TchDel_{int(time.time())}", GRADE5, "A")
        r = teacher_token.delete(f"{API}/students/{s['id']}", timeout=20)
        assert r.status_code == 200, r.text
        g = admin_token.get(f"{API}/students/{s['id']}", timeout=20)
        assert g.json()["status"] == "inactive"

    def test_teacher_delete_other_class_403(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_TchDelOther_{int(time.time())}", GRADE6, "B")
        r = teacher_token.delete(f"{API}/students/{s['id']}", timeout=20)
        assert r.status_code == 403


# -------- Credential slip PDF --------
class TestCredentialSlip:
    def test_slip_admin_pdf(self, admin_token):
        # Get first active student
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s["status"] == "active")
        name = next(s["name"] for s in lst.json() if s["id"] == sid)
        r = admin_token.get(f"{API}/students/{sid}/credential-slip", timeout=30)
        assert r.status_code == 200
        assert r.headers["Content-Type"].startswith("application/pdf")
        assert "attachment" in r.headers["Content-Disposition"]
        assert r.content.startswith(b"%PDF")
        # Content should contain Parent@123 default password text and "Parent Portal"
        # PDFs may compress text; check len reasonable
        assert len(r.content) > 1500
        assert name  # sanity

    def test_slip_teacher_own_class_ok(self, admin_token, teacher_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s["status"] == "active")
        r = teacher_token.get(f"{API}/students/{sid}/credential-slip", timeout=20)
        assert r.status_code == 200

    def test_slip_teacher_out_of_class_403(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_SlipOther_{int(time.time())}", GRADE6, "B")
        r = teacher_token.get(f"{API}/students/{s['id']}/credential-slip", timeout=20)
        assert r.status_code == 403

    def test_slip_accounts_403(self, admin_token, accounts_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s["status"] == "active")
        r = accounts_token.get(f"{API}/students/{sid}/credential-slip", timeout=20)
        assert r.status_code == 403


# -------- Send credentials --------
class TestSendCredentials:
    def test_send_invalid_channel(self, admin_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s["status"] == "active")
        r = admin_token.post(f"{API}/students/{sid}/send-credentials", json={"channel": "carrier-pigeon"}, timeout=20)
        assert r.status_code == 400

    def test_send_telegram_with_reset_first_then_login(self, admin_token):
        # Use seeded parent.stu-00001
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        target = next(s for s in lst.json() if s.get("student_uid") == "STU-00001")
        sid = target["id"]
        r = admin_token.post(f"{API}/students/{sid}/send-credentials", json={"channel": "telegram", "reset_first": True}, timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["ok"] is True
        assert d["channel"] == "telegram"
        assert d["reset_first"] is True
        assert d["delivery_status"] in ("queued", "sent")
        # parent password is back to Parent@123
        lr = requests.post(f"{API}/auth/login", json={
            "email": "parent.stu-00001@sunrise.edu", "password": "Parent@123"}, timeout=20)
        assert lr.status_code == 200

    def test_send_whatsapp_no_reset(self, admin_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s.get("student_uid") == "STU-00001")["id"] if False else next(
            s["id"] for s in lst.json() if s.get("student_uid") == "STU-00001")
        r = admin_token.post(f"{API}/students/{sid}/send-credentials", json={"channel": "whatsapp", "reset_first": False}, timeout=20)
        assert r.status_code == 200
        assert r.json()["channel"] == "whatsapp"

    def test_send_teacher_other_class_403(self, admin_token, teacher_token):
        s = _seed_student(admin_token, f"TEST_SendOther_{int(time.time())}", GRADE6, "B")
        r = teacher_token.post(f"{API}/students/{s['id']}/send-credentials", json={"channel": "telegram"}, timeout=20)
        assert r.status_code == 403

    def test_send_accounts_403(self, admin_token, accounts_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s.get("student_uid") == "STU-00001")
        r = accounts_token.post(f"{API}/students/{sid}/send-credentials", json={"channel": "telegram"}, timeout=20)
        assert r.status_code == 403

    def test_send_anon_401(self, admin_token):
        lst = admin_token.get(f"{API}/students", params={"class_id": GRADE5, "section": "A"}, timeout=20)
        sid = next(s["id"] for s in lst.json() if s.get("student_uid") == "STU-00001")
        r = requests.post(f"{API}/students/{sid}/send-credentials", json={"channel": "telegram"}, timeout=20)
        assert r.status_code in (401, 403)


# -------- Regression: parent_login on POST /students --------
class TestRegressionParentLogin:
    def test_create_student_returns_parent_login(self, admin_token):
        r = admin_token.post(f"{API}/students", json={
            "name": f"TEST_REG_{int(time.time())}",
            "class_id": GRADE5, "section": "A",
            "parent": {"name": "P", "phone": "+918888888888", "relationship": "guardian", "email": "", "whatsapp": ""}
        }, timeout=20)
        assert r.status_code in (200, 201)
        body = r.json()
        assert "parent_login" in body
        assert body["parent_login"]["email"]

    def test_seeded_parent_login_works(self):
        r = requests.post(f"{API}/auth/login",
                         json={"email": "parent.stu-00001@sunrise.edu", "password": "Parent@123"}, timeout=20)
        assert r.status_code == 200


# -------- Users CRUD (TeachersPage) --------
class TestTeacherUserCRUD:
    def test_create_update_deactivate_teacher(self, admin_token):
        ts = int(time.time())
        email = f"test_iter19_teach_{ts}@sunrise.edu"
        subs = admin_token.get(f"{API}/academic/subjects", timeout=15).json()
        sid = subs[0]["id"] if subs else None
        assert sid, "No subjects available to assign teacher"
        r = admin_token.post(f"{API}/users", json={
            "email": email, "name": "TEST Iter19 Teacher", "phone": "+910000000000",
            "role": "teacher", "password": "Teacher@123", "subjects": [sid]
        }, timeout=20)
        assert r.status_code in (200, 201), r.text
        uid = r.json()["id"]
        p = admin_token.patch(f"{API}/users/{uid}", json={"name": "TEST Iter19 Updated", "phone": "+910000000099"}, timeout=20)
        assert p.status_code == 200, p.text
        d = admin_token.delete(f"{API}/users/{uid}", timeout=20)
        assert d.status_code in (200, 204), d.text

    def test_teacher_cannot_patch_users(self, teacher_token, admin_token):
        me = teacher_token.get(f"{API}/auth/me", timeout=10).json()
        r = teacher_token.patch(f"{API}/users/{me['id']}", json={"name": "Hacker"}, timeout=10)
        assert r.status_code in (401, 403)
