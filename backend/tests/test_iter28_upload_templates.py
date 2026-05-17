"""Tests for Excel/CSV bulk upload templates with data validations (iter 28).

Verifies:
- /api/uploads/template/{kind}.xlsx: 2 sheets (Upload + References), DataValidation
  attached to the correct columns and references Sunrise's real data.
- /api/uploads/template-csv/{kind}: returns CSV with headers + hint row.
- RBAC: school_admin/teacher allowed, parent gets 403.
- Hint row parser-skip: uploading a freshly-downloaded xlsx -> 0 inserted, 0 errors.
- Uploading a template with one real row -> 1 inserted, 0 errors.
"""
import io
import os
import openpyxl
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL"):
                BASE_URL = line.split("=", 1)[1].strip().strip('"').rstrip("/")
API = f"{BASE_URL}/api"


def _login(email, password):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=20)
    assert r.status_code == 200, f"login failed for {email}: {r.status_code} {r.text}"
    return s


@pytest.fixture(scope="module")
def admin():
    return _login("principal@sunrise.edu", "School@123")


@pytest.fixture(scope="module")
def teacher():
    return _login("priya@sunrise.edu", "Teacher@123")


@pytest.fixture(scope="module")
def parent():
    return _login("parent.stu-00001@sunrise.edu", "Parent@123")


def _download_xlsx(session, kind):
    r = session.get(f"{API}/uploads/template/{kind}", timeout=30)
    assert r.status_code == 200, f"{kind}: {r.status_code} {r.text[:200]}"
    assert "spreadsheetml" in r.headers.get("content-type", "")
    return openpyxl.load_workbook(io.BytesIO(r.content))


# ----- xlsx structural tests -----
class TestXlsxTemplates:
    def test_students_template(self, admin):
        wb = _download_xlsx(admin, "students")
        assert "Upload" in wb.sheetnames
        assert "References" in wb.sheetnames
        ws = wb["Upload"]
        headers_row = [c.value for c in ws[1]]
        assert headers_row[1] == "class_id"
        dvs = list(ws.data_validations.dataValidation)
        assert len(dvs) >= 1
        joined = " ".join([str(dv.formula1) + " " + str(list(dv.sqref)) for dv in dvs])
        assert "References" in joined
        assert "B3:B" in joined  # class_id col
        ref = wb["References"]
        names = " ".join(str(ref.cell(row=r, column=2).value)
                         for r in range(1, ref.max_row + 1)
                         if ref.cell(row=r, column=2).value)
        assert "Grade 5" in names and "Grade 6" in names

    def test_attendance_template(self, admin):
        wb = _download_xlsx(admin, "attendance")
        assert "References" in wb.sheetnames
        ws = wb["Upload"]
        dvs = list(ws.data_validations.dataValidation)
        targets = " ".join(str(list(dv.sqref)) for dv in dvs)
        assert "A3:A" in targets  # student_uid
        assert "C3:C" in targets  # status
        ref = wb["References"]
        ids_text = " ".join(str(ref.cell(row=r, column=1).value)
                            for r in range(1, ref.max_row + 1)
                            if ref.cell(row=r, column=1).value)
        for s in ["present", "absent", "sick", "late", "half_day", "excused"]:
            assert s in ids_text, f"missing status {s}"
        for i in range(1, 9):
            assert f"STU-{i:05d}" in ids_text, f"missing STU-{i:05d}"

    def test_marks_template(self, admin):
        wb = _download_xlsx(admin, "marks")
        ws = wb["Upload"]
        dvs = list(ws.data_validations.dataValidation)
        targets = " ".join(str(list(dv.sqref)) for dv in dvs)
        for col in ["A3:A", "B3:B", "C3:C"]:
            assert col in targets, f"missing DV for {col}"
        ref = wb["References"]
        text = " ".join(str(ref.cell(row=r, column=c).value)
                        for r in range(1, ref.max_row + 1) for c in (1, 2)
                        if ref.cell(row=r, column=c).value)
        assert "STU-00008" in text

    def test_holidays_template(self, admin):
        wb = _download_xlsx(admin, "holidays")
        ws = wb["Upload"]
        headers_row = [c.value for c in ws[1]]
        type_idx = headers_row.index("type") + 1
        col_letter = openpyxl.utils.get_column_letter(type_idx)
        dvs = list(ws.data_validations.dataValidation)
        targets = " ".join(str(list(dv.sqref)) for dv in dvs)
        assert f"{col_letter}3:{col_letter}" in targets
        ref = wb["References"]
        ids_text = " ".join(str(ref.cell(row=r, column=1).value)
                            for r in range(1, ref.max_row + 1)
                            if ref.cell(row=r, column=1).value)
        for t in ["public", "school", "religious", "festival"]:
            assert t in ids_text


# ----- CSV templates -----
class TestCsvTemplates:
    @pytest.mark.parametrize("kind", ["students", "attendance", "marks", "holidays"])
    def test_csv_template(self, admin, kind):
        r = admin.get(f"{API}/uploads/template-csv/{kind}", timeout=20)
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        body = r.content.decode("utf-8").splitlines()
        assert len(body) >= 2
        assert "e.g." in body[1].lower() or "/" in body[1].lower()


# ----- RBAC -----
class TestRBAC:
    def test_parent_forbidden(self, parent):
        r = parent.get(f"{API}/uploads/template/students", timeout=15)
        assert r.status_code == 403

    def test_teacher_allowed(self, teacher):
        r = teacher.get(f"{API}/uploads/template/students", timeout=15)
        assert r.status_code == 200

    def test_unknown_kind_returns_404(self, admin):
        r = admin.get(f"{API}/uploads/template/foobar", timeout=15)
        assert r.status_code == 404


# ----- Hint row skip behaviour -----
class TestHintRowSkipped:
    def test_empty_template_upload_inserts_zero(self, admin):
        r = admin.get(f"{API}/uploads/template/holidays", timeout=20)
        assert r.status_code == 200
        files = {"file": ("holidays.xlsx", r.content,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        up = admin.post(f"{API}/uploads/holidays", files=files, timeout=30)
        assert up.status_code == 200, up.text
        data = up.json()
        assert data["inserted"] == 0
        assert len(data["errors"]) == 0

    def test_template_with_one_real_row(self, admin):
        r = admin.get(f"{API}/uploads/template/holidays", timeout=20)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb["Upload"]
        ws.append(["TEST_iter28_holiday", "2099-12-25", "2099-12-25", "school", "test row"])
        buf = io.BytesIO()
        wb.save(buf)
        files = {"file": ("holidays.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        up = admin.post(f"{API}/uploads/holidays", files=files, timeout=30)
        assert up.status_code == 200, up.text
        data = up.json()
        assert data["inserted"] == 1, f"expected 1, got {data}"
        assert len(data["errors"]) == 0
